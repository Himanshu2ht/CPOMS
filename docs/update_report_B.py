"""Batch B: Chapters 4-5 (requirements + analysis)."""
from docx import Document

P = "docs/CPOMS_Project_Report.docx"
d = Document(P)
paras = d.paragraphs


def sub_para(i, old, new):
    p = paras[i]
    assert old in p.text, f"[para {i}] not found: {old[:60]!r}"
    p.text = p.text.replace(old, new)


def set_para(i, text):
    paras[i].text = text


def row_set_by_firstcol(ti, firstcol, col, new):
    t = d.tables[ti]
    for r in t.rows[1:]:
        if r.cells[0].text.strip().startswith(firstcol):
            r.cells[col].text = new
            return
    raise AssertionError(f"table {ti} has no row starting {firstcol!r}")


# ---- 4.1.1 screens ----
sub_para(105, "Customers see Home/Menu, Cart, Checkout-Pay, My Orders (live status + QR), and Feedback.",
         "Customers see Menu, Cart, My Orders and Order detail (live status badge, timeline, slot countdown and pickup token).")
sub_para(105, "Staff see KDS Queue (auto-refresh, large Preparing/Ready buttons, new-order chime) and Collection Scan (camera input + manual token entry). The Manager sees Menu/Slot Publisher, Order Monitor, Refunds, and Reports.",
         "Staff see the Kitchen Dashboard and Queue (auto-refresh, per-order countdowns, one-tap advance buttons, token handover box). Admins see the Admin Dashboard, Menu & Stock, Time Slots, Users & Staff, All Orders and Reports (sales, top items, no-show log).")

# ---- 4.1.3 software interfaces ----
set_para(111, "Frontend: HTML5, CSS3, Bootstrap 5.3 CDN, vanilla JavaScript (fetch polling — every 8 s on the order page, every 10 s in the kitchen). Backend: Python 3.11+, Flask 3.1, Jinja2 templates, Flask-SQLAlchemy 3.1, Flask-Login 0.6 with strong session protection, Werkzeug scrypt hashing, Gunicorn in production. Database: MySQL 8 locally (modelled in MySQL Workbench), PostgreSQL on Supabase once deployed, SQLite as a local fallback; SQLAlchemy create_all builds the schema and an empty database seeds itself on first boot. External: a sandbox payment adapter with HMAC-signed webhooks stands in for a live gateway (UPI/cards) — swap create_payment() for the real SDK to go live. All secrets live in environment variables, never in code or Git. Times are stored as naive UTC datetimes; money as DECIMAL rupees.")

# ---- 4.1.4 comms ----
set_para(113, "All browser traffic uses HTTPS (TLS 1.2+); lab deployment uses a self-signed cert with a documented exception, production uses Let's Encrypt. The sandbox payment adapter signs its callbacks with HMAC and the app verifies every signature before trusting it; duplicate callbacks are acknowledged without side effects. Status polling uses lightweight JSON — order pages every 8 seconds, the kitchen every 10 — with no WebSockets, which keeps lab and free-tier hosting simple. Notifications are on-screen (flash messages, live badges, timelines, countdowns); no email is sent in this release, so nothing ever blocks order confirmation.")

# ---- enforcement detail ----
set_para(120, "Enforcement detail — how online-only is guaranteed: UI shows no cash option and only the sandbox checkout path exists; routes advance orders strictly through legal status transitions; the service layer issues the pickup token only after payment success; duplicate webhooks return success without side effects; and background sweeps expire unpaid orders and stale Ready orders on their own. UI, routes, services and sweeps each enforce it — so testing it once is never enough; Chapter 9 tests it several ways.")

# ---- 4.5 attributes ----
sub_para(128, "session timeout (30 min idle),", "sessions that end with the browser and strong session protection,")
sub_para(128, "and Pytest coverage ≥70% on payment/order logic.", "and a 19-test Pytest suite covering auth, ordering, payments and the kitchen flow.")
sub_para(128, "Auditability: immutable logs for every money or status event, exportable for the manager and examiners.",
         "Auditability: every order and payment row keeps its status and timestamps, exportable via the reports pages for the manager and examiners.")
sub_para(128, "Portability: pure-Python + MySQL,", "Portability: pure-Python + MySQL/PostgreSQL/SQLite,")

# ---- FR table 16 (col 1 = Requirement) ----
fr = {
 "FR-01": "Register with name, email and a strong password (8+ chars with upper, lower, digit and special); duplicate emails rejected with a clear message.",
 "FR-02": "Login/logout with Flask-Login (strong session protection); 5 wrong attempts from one IP+email lock logins for 5 minutes; passwords stored only as Werkzeug scrypt hashes.",
 "FR-03": "Browse menu by category; see price, live stock and availability; instant search without reload.",
 "FR-04": "Manager adds items (name, category, price, stock); edits and availability toggles apply immediately; categories can be added with or without an item.",
 "FR-05": "Define daily slots (time + cap, e.g., 80); overbooking blocked; seats-left shown; past slots roll forward automatically.",
 "FR-06": "Cart: add/remove, qty capped at 10 per line; unpaid orders lapse after 10 minutes; slot countdown shown.",
 "FR-08": "Sandbox gateway adapter: HMAC-signed callbacks verified before trusting; duplicate callbacks ignored (idempotent).",
 "FR-09": "Issue the text pickup token ONLY after Paid; unpaid orders never get one.",
 "FR-10": "Live tracking across Pending Payment, Paid, Preparing, Ready, Collected (plus Cancelled/Failed/No-Show); order page refreshes every 8 s.",
 "FR-11": "Kitchen queue sorted by slot then order time; auto-refresh every 10 s, with action buttons surviving every refresh.",
 "FR-12": "Staff advance Paid to Preparing to Ready to Collected; illegal jumps rejected; token handover or No-Show closes Ready orders.",
 "FR-13": "Token box capitalises input itself; handover allowed only for Ready orders; wrong tokens rejected with a message.",
 "FR-14": "Cancel free while unpaid; Paid orders refundable within 5 minutes of ordering (sandbox-simulated refund, stock restored).",
 "FR-15": "Unavailable or out-of-stock items are blocked at checkout with a clear message; cancelling restores stock and frees the slot.",
 "FR-16": "Reports: daily sales, top items and the no-show/waste log for the manager.",
 "FR-17": "No-show handling: Ready orders uncollected 15 minutes past slot close as No-Show (no refund) into the waste log.",
 "FR-18": "Live badges, timelines and countdowns update in-page on every status change; flash messages confirm each action.",
 "FR-19": "Admin Users & Staff page: accounts grouped by role (admin/staff/customer); change roles, delete order-free accounts; all-orders view included.",
 "FR-21": "Every order and payment keeps its status history with timestamps for reconciliation.",
}
for fid, text in fr.items():
    row_set_by_firstcol(16, fid, 1, text)

# ---- PR table 17 ----
row_set_by_firstcol(17, "PR-01", 0, "PR-01 | Menu page loads ≤1.5 s on lab Wi-Fi; checkout ≤2 s (excluding bank page).")
row_set_by_firstcol(17, "PR-03", 0, "PR-03 | Kitchen refresh every 10 s; order page every 8 s; Ready visible to the customer on the next poll.")
row_set_by_firstcol(17, "PR-05", 0, "PR-05 | Day-end report (2000 orders) renders ≤5 s.")
# rows are single-column "ID | text"; handle combined cell
for r in d.tables[17].rows[1:]:
    c = r.cells[0].text.strip()
    if c.startswith("PR-01"):
        r.cells[0].text = "PR-01 | Menu page loads ≤1.5 s on lab Wi-Fi; checkout ≤2 s (excluding bank page)."
    elif c.startswith("PR-03"):
        r.cells[0].text = "PR-03 | Kitchen refresh every 10 s; order page every 8 s; Ready visible to the customer on the next poll."
    elif c.startswith("PR-05"):
        r.cells[0].text = "PR-05 | Day-end report (2000 orders) renders ≤5 s."

# ---- DC table 18 ----
for r in d.tables[18].rows[1:]:
    c = r.cells[0].text.strip()
    if c.startswith("DC-02"):
        r.cells[0].text = "DC-02 | Money as DECIMAL rupees; datetimes naive UTC; passwords never plain (Werkzeug scrypt hash)."
    elif c.startswith("DC-04"):
        r.cells[0].text = "DC-04 | SQLAlchemy create_all is the runtime source of truth; empty databases self-seed demo data on boot."

# ---- UC-08 cancel rule ----
sub_para(140, "allowed only if ≥30 min before slot start and status is Paid (not yet Preparing). System calls gateway Refund API with transaction ID, marks order Refunded on success, notifies customer, and restores stock/slot capacity. Refunds never disburse cash; the gateway settles to the original UPI/card. Manager-initiated refunds (item unavailable) follow the same API path with an approval timestamp for audit.",
         "allowed while the order is still unpaid, or within 5 minutes of ordering for Paid orders (before the kitchen starts). Cancelling restores stock, frees the slot and flags the sandbox payment refunded; Ready orders past their grace period are instead closed as No-Show. Refunds never disburse cash — a live gateway would settle to the original UPI/card.")

# ---- L1 process descriptions ----
set_para(147, "Level-1 process descriptions (inputs → outputs): 1.0 Auth checks the email and Werkzeug scrypt hash against D1 Users and opens a Flask-Login session (5 wrong tries from one IP+email means a 5-minute lockout). 2.0 Menu/Slot Management lets the manager write D2 Menu and D3 Slots with caps; categories live in their own table. 3.0 Cart & Checkout checks D2 stock and D3 capacity plus the 25-minute prep lead, then writes a PENDING_PAYMENT order to D4 with a 10-minute expiry sweep. 4.0 Online Payment is the gate: it starts a sandbox payment, verifies the HMAC callback, enforces idempotency on the gateway reference in D5 Payments, and promotes the order to Paid exactly once. 5.0 Tracking reads D4/D5 and renders status, countdown and token (the token is only ever issued for Paid). 6.0 Kitchen & Collection lists Paid by slot/time from D4, records Preparing/Ready/Collected/No-Show with a typed token at handover. 7.0 Reports aggregates D4/D5 into sales, top items and the no-show log. Every flow conserves data: stock drops only on Paid and returns on cancel; slot bookings are freed on expiry, cancel or failure.")

# ---- dictionary intro + ER preview ----
sub_para(149, "Amounts are integers (paise); statuses are enums;",
         "Amounts are DECIMAL rupees; statuses are enums;")
set_para(157, "Seven tables implement the dictionary: users, menu_items, categories, time_slots, orders, order_items, payments. Users 1—N orders; time_slots 1—N orders; orders 1—N order_items; orders 1—1 payments (gateway reference unique); categories 1—N menu items. The service layer — not the database — enforces the online-only rule (a token exists only on Paid rows), which is why Chapter 9 tests behaviour rather than constraints. The full schema with column types and keys is detailed in Chapter 7; what matters here is that the ER covers every DFD store with no extra cash-related entity — because none exists.")

# ---- UC catalogue table 20 ----
uc = {
 "UC-03": "UC-03 | Cart + Slot | Customer | Items + slot → unpaid order (10-min expiry)",
 "UC-05": "UC-05 | Track + token | Customer | Paid → token + live status",
 "UC-06": "UC-06 | Prepare + Collect | Staff/Desk | Paid → Collected via queue + token",
 "UC-08": "UC-08 | Cancel/Refund | Cust + PG | Unpaid/≤5-min Paid → Cancelled (+sandbox refund)",
 "UC-09": "UC-09 | Reports | Manager | Orders → sales + no-show log",
 "UC-10": "UC-10 | No-show sweep | System | Stale Ready → No-Show log",
}
for r in d.tables[20].rows[1:]:
    cells = [c.text.strip() for c in r.cells]
    if cells and cells[0] in uc:
        r.cells[0].text = uc[cells[0]].split(" | ")[0]
        # rewrite whole row
        parts = uc[cells[0]].split(" | ")
        for ci, val in enumerate(parts):
            r.cells[ci].text = val

# ---- flows table 23 ----
flows = [
 ("Menu_Item", "Menu_Item | item_id + name + price DECIMAL (>0) + category (dropdown) + stock (≥0) + active flag"),
 ("Slot", "Slot | slot_id + daily time (HH:MM, no date) + cap (e.g., 80) + booked (≤cap); past slots roll forward"),
 ("Cart_Hold", "Pending_Order | user_id + slot_id + lines(item,qty ≤10) + PENDING_PAYMENT + 10-min expiry sweep"),
 ("Gateway_Order", "Gateway_Order | gateway ref + amount (= order total) + hosted checkout URL"),
 ("Webhook_Callback", "Webhook_Callback | gateway ref + success flag + HMAC signature (must verify); duplicates ignored"),
 ("Order (D4)", "Order (D4) | order_id + user_id + slot_id + total DECIMAL + status ∈ {Pending Payment, Paid, Preparing, Ready, Collected, Cancelled, Failed, No-Show} + token (Paid only)"),
 ("Payment (D5)", "Payment (D5) | payment_id + order_id (FK) + gateway ref (UNIQUE) + amount + status ∈ {Created, Success, Failed, Refunded}"),
 ("Feedback (D6)", "NoShow_Log | order_id (FK, Ready past grace) + slot + total; no refund, kept as waste record"),
]
for r in d.tables[23].rows[1:]:
    first = r.cells[0].text.strip().split(" | ")[0].split()[0]
    for key, full in flows:
        if r.cells[0].text.strip().startswith(key):
            parts = full.split(" | ")
            r.cells[0].text = parts[0]
            r.cells[1].text = parts[1]
            break

# ---- traceability table 38 ----
for r in d.tables[38].rows[1:]:
    t = " | ".join(c.text.strip() for c in r.cells)
    if t.startswith("FR-06"):
        r.cells[2].text = "SEQ-01 / place_order"
    elif t.startswith("FR-09"):
        r.cells[2].text = "models.Order.token"
        r.cells[3].text = "TC-09"
    elif t.startswith("FR-11/12"):
        r.cells[2].text = "SEQ-02 / staff_bp"
    elif t.startswith("FR-13"):
        r.cells[2].text = "SEQ-02 / token verify"

d.save(P)
print("Batch B saved.")
