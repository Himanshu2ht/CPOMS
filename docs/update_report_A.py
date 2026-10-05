"""Batch A: front matter + Chapters 1-3. Run from repo root: .venv/bin/python docs/update_report_A.py"""
from docx import Document

P = "docs/CPOMS_Project_Report.docx"
d = Document(P)
paras = d.paragraphs


def sub_para(i, old, new):
    p = paras[i]
    assert old in p.text, f"[para {i}] not found: {old[:60]!r}"
    p.text = p.text.replace(old, new)


def row_set_by_firstcol(ti, firstcol, col, new):
    t = d.tables[ti]
    for r in t.rows[1:]:
        if r.cells[0].text.strip() == firstcol:
            r.cells[col].text = new
            return
    raise AssertionError(f"table {ti} has no row starting {firstcol!r}")


def cell_sub(ti, old, new):
    t = d.tables[ti]
    found = False
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                if old in p.text:
                    p.text = p.text.replace(old, new)
                    found = True
    assert found, f"table {ti} missing: {old[:60]!r}"


# ---- cover ----
sub_para(8, "Failed payments auto-cancel the cart hold within 15 minutes.",
         "Failed or abandoned payments expire on their own after 10 minutes, freeing the slot for others.")
sub_para(9, "DB: MySQL Community Server (via MySQL Workbench)  |  Testing: Pytest  |  Tools: Git, GitHub, VS Code.",
         "DB: MySQL locally, PostgreSQL (Supabase) live, SQLite fallback  |  Hosting: Render + Gunicorn  |  Testing: Pytest (19 tests)  |  Tools: Git, GitHub, VS Code.")

# ---- abstract ----
sub_para(26, "(Placed → Paid → Preparing → Ready → Collected)",
         "(Pending Payment → Paid → Preparing → Ready → Collected, with Cancelled, Failed and No-Show branches)")
sub_para(27, "MySQL Community Server managed via MySQL Workbench; Pytest for testing",
         "MySQL for local work and PostgreSQL (Supabase) in production, deployed on Render with Gunicorn; Pytest (19 automated tests) for testing")

# ---- ch1 ----
sub_para(39, "All three share a single MySQL database accessed via Flask-SQLAlchemy",
         "All three share one database — MySQL on a developer machine, PostgreSQL (Supabase) once deployed — accessed through Flask-SQLAlchemy")
sub_para(40, "Students log in with college email, add items to a cart, choose a slot, and pay online.",
         "Students log in with their college email, add items to a cart, pick a daily pickup time, and pay online.")
sub_para(40, "The gateway returns a transaction ID; only then is the order marked Paid and queued for the kitchen in slot order.",
         "Only after the gateway confirms payment is the order marked Paid, given a short pickup token (like T0012), and queued for the kitchen in slot order.")
sub_para(40, "The kitchen marks items Preparing and Ready; the student receives live updates and a QR/token code.",
         "The kitchen moves it through Preparing to Ready while the student watches live status and a slot countdown.")
sub_para(40, "Failed or expired payments release the held stock automatically so that seats and items are never blocked by unpaid carts.",
         "If payment fails or is abandoned, the unpaid order lapses after 10 minutes and its slot booking is released automatically.")
sub_para(48, "Order means a paid order unless explicitly called Cart (unpaid).",
         "Cart means the pre-checkout selection kept in the browser session; Order means a saved order row, paid or not.")
sub_para(48, "Status values are always exactly one of Placed, Paid, Preparing, Ready, Collected, Cancelled, Refunded — no synonyms are used in code or UI.",
         "Status is always exactly one of Pending Payment, Paid, Preparing, Ready, Collected, Cancelled, Failed or No-Show (a Ready order still uncollected 15 minutes past its slot) — the same words appear in code, UI and tests.")

# ---- ch2 ----
sub_para(64, "timestamped in MySQL", "timestamped in the database")
sub_para(65, "holds cart stock for only 15 minutes pending payment, and never generates a pickup QR without a successful gateway callback.",
         "lets unpaid orders lapse after 10 minutes, and never issues a pickup token without a successful gateway callback.")

# ---- ch3 ----
sub_para(81, "The system runs as a Python Flask app on a lab/department server, serving Jinja2-rendered Bootstrap pages, persisting to MySQL 8 via SQLAlchemy, and calling the gateway over HTTPS with webhook callbacks.",
         "The system runs as a Python Flask app — on a lab machine during development and on Render (Gunicorn) in production — serving Jinja2-rendered Bootstrap pages, persisting to MySQL 8 locally or PostgreSQL (Supabase) once deployed, and confirming payment through HMAC-signed gateway webhooks.")
sub_para(84, "the collection desk needs a device (tablet/phone) to scan QR codes even if the printer fails.",
         "the collection desk needs a device (tablet/phone) where staff type the pickup token into a box that capitalises on its own — handover happens only for Ready orders.")
sub_para(86, "Menu management lets the manager publish daily items with price, photo, veg/non-veg tags, availability counts, and slot associations.",
         "Menu management lets the manager publish items with price, category, stock and availability, and add new categories on their own.")
sub_para(86, "Slot management defines pickup windows and per-slot capacity caps that prevent overbooking.",
         "Slots are daily pickup times with per-slot capacity caps that prevent overbooking; past slots roll forward by themselves so the list never goes stale.")
sub_para(86, "Cart and checkout hold stock for 15 minutes while the customer pays online — expiry releases the hold automatically.",
         "Checkout creates an unpaid order first — if nobody pays within 10 minutes it lapses and the slot booking is released automatically.")
sub_para(87, "Payment processing is the heart: it creates a gateway order, redirects to UPI/card/net-banking, verifies the callback signature, marks the order Paid exactly once (idempotent), and never confirms without success. Order tracking shows live status and QR pickup code only for paid orders. Staff operations provide the KDS queue (sorted by slot then time), Preparing/Ready transitions, collection scan, and exception handling (item unavailable → refund-to-source). Finally, reporting and feedback give sales summaries, slot utilisation, and ratings that feed tomorrow's menu planning.",
         "Payment is the heart: a sandbox gateway adapter starts the payment, the customer pays, the signed callback is verified, and the order turns Paid exactly once — even if the callback arrives twice. Stock drops only at that moment, and only then is the pickup token issued. Tracking shows live status, a slot countdown and the token. Staff get a kitchen dashboard and queue that refresh every 10 seconds, one-tap Preparing/Ready buttons, token handover, and manual or automatic no-show marking. Admins get revenue and order stats, menu and slot management, a Users & Staff page for roles, an all-orders view, and a no-show log that doubles as the waste record.")
sub_para(93, "uses auto-refresh every 15 seconds with an audible cue.",
         "uses auto-refresh every 10 seconds without any page reload.")
sub_para(95, "Technical: must run on Flask 3 + MySQL 8 in the lab;",
         "Technical: must run on Flask 3 with MySQL 8 in the lab and PostgreSQL (Supabase) in production, served by Gunicorn on Render;")
sub_para(95, "must expire unpaid holds in 15 minutes,",
         "must let unpaid orders lapse after 10 minutes,")
sub_para(95, "menu edits after slot start apply only to future slots.",
         "slot times repeat daily and roll forward automatically.")
sub_para(99, "share one MySQL truth,", "share one database,")

# ---- glossary table 5 ----
row_set_by_firstcol(5, "Slot", 1, "A daily collection time (e.g., 12:30) with a capacity cap; identical every day, no dates.")
row_set_by_firstcol(5, "Token / QR", 1, "Short text pickup code (e.g., T0012) issued only after payment success; typed at handover.")
t5 = d.tables[5]
for name, meaning in [
    ("No-Show", "A Ready order still uncollected 15 minutes past its slot; closed without refund and kept in the waste log."),
    ("Grace window", "15 minutes after slot start during which a Ready order can still be collected."),
]:
    row = t5.add_row()
    row.cells[0].text = name
    row.cells[1].text = meaning

# ---- function groups table 11 ----
cell_sub(11, "15-min hold; slot chosen", "10-min unpaid expiry; slot chosen")
row_set_by_firstcol(11, "F6 Tracking & QR", 0, "F6 Tracking & token")
row_set_by_firstcol(11, "F6 Tracking & token", 2, "Live status + countdown + pickup code")
row_set_by_firstcol(11, "F8 Reports & Feedback", 0, "F8 Reports & no-show log")
row_set_by_firstcol(11, "F8 Reports & no-show log", 2, "Sales + waste record")

d.save(P)
print("Batch A saved.")
