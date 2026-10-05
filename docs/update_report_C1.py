"""Batch C part 1: Chapters 7-8 (design + code)."""
from docx import Document

P = "docs/CPOMS_Project_Report.docx"
d = Document(P)
paras = d.paragraphs


def sub_para(i, old, new):
    p = paras[i]
    assert old in p.text, f"[para {i}] not found: {old[:60]!r}"
    p.text = p.text.replace(old, new)


def set_para(i, text):
    paras[i].text = text


def row_set_by_firstcol(ti, firstcol, col, new):
    t = d.tables[ti]
    for r in t.rows[1:]:
        if r.cells[0].text.strip().startswith(firstcol):
            r.cells[col].text = new
            return
    raise AssertionError(f"table {ti} has no row starting {firstcol!r}")


# ---- 7.1 lifecycle ----
set_para(193, "Request lifecycle (example checkout): the browser POSTs /checkout with slot_id → orders_bp calls order_service.place_order(), which checks the cart, the slot (exists, has room, starts at least 25 minutes out) and the stock → a PENDING_PAYMENT order plus its lines is saved and the slot booked → the sandbox gateway answers with a checkout URL → the customer pays (or simulates failure) → the app applies payment_service.handle_event(), which verifies the reference, shrugs off duplicates, marks Paid, issues token T0001 and drops the stock — all in one commit → staff see it on the next 10-second queue poll. Failures never promote: unknown references are rejected, duplicates return success without side effects, and FAILED orders free their slot. This single path is why duplicate charges and half-paid orders are structurally impossible.")

# ---- module table 33 ----
mods = {
 "Auth": ("Register/login, scrypt hash, 5-fail lockout, roles, profile page", "auth blueprint; models.User; login/profile pages"),
 "Menu/Slot Admin": ("Menu + standalone categories + daily slots, stock control", "admin blueprint (/admin/menu, /admin/slots, /admin/users)"),
 "Cart & Hold": ("Cart + unpaid order, 10-min expiry sweep, slot countdown", "orders blueprint; order_service.place_order()"),
 "Payment (critical)": ("Sandbox adapter, HMAC verify, idempotent confirm, sandbox refund", "payments blueprint (/payments/*, /webhook); payment_service"),
 "Tracking/QR": ("Status + countdown + token, Paid only", "orders/detail page; /orders/<id>/status JSON"),
 "KDS + Collection": ("Live queue, advance buttons, token handover, no-show", "staff blueprint (/staff, /staff/queue); queue-data JSON"),
 "Reports/Feedback": ("Sales, top items, no-show log, all-orders view", "admin blueprint (/admin, /reports, /orders)"),
}
for first, (resp, files) in mods.items():
    row_set_by_firstcol(33, first, 1, resp)
    row_set_by_firstcol(33, first, 2, files)

# ---- 7.2 heading + design paras ----
sub_para(196, "7.2 Database Design (MySQL Workbench → 8 Tables)", "7.2 Database Design (MySQL Workbench → 7 Tables)")
set_para(197, "Designed in MySQL Workbench (EER) and then implemented as SQLAlchemy models — Chapter 8 shows the real code. Money is DECIMAL rupees; times are naive UTC datetimes; statuses are a Python enum stored per row; passwords are scrypt hashes only. Foreign keys link orders to users and slots, order lines to orders, and payments to orders. Uniqueness is enforced on users.email (login), payments.gateway_ref (the idempotency anchor) and orders.token (one token per paid order).")
set_para(200, "Indexing and integrity: a unique index on users.email for login, an index on orders.user_id for history, and slot lookups by id for the kitchen queue. A background sweep on every request expires unpaid orders after 10 minutes and closes stale Ready orders as No-Show; both settle their slot bookings. Backups come from Supabase in production (mysqldump in the lab); a restore is drilled before UAT so a failed demo never loses the menu.")

# ---- schema table 34: 7 real tables ----
t34 = d.tables[34]
new_rows = [
 ("users", "id PK, name, email UNIQUE + index, password_hash (scrypt), role", "Plaintext never stored; roles assigned by admin only"),
 ("menu_items", "id PK, name, category, price NUMERIC(8,2), stock, is_available", "Unavailable or out-of-stock lines blocked at checkout"),
 ("categories", "id PK, name UNIQUE", "Standalone list; addable with or without any item"),
 ("time_slots", "id PK, start_time, capacity, booked", "Daily times, no dates; past slots roll forward by themselves"),
 ("orders", "id PK, user FK, slot FK, total NUMERIC(10,2), status ENUM(8), token UNIQUE, created_at", "Token set only on Paid; history indexed by user"),
 ("order_items", "id PK, order FK, item FK, quantity, unit_price", "Cascade-delete with the order"),
 ("payments", "id PK, order FK, gateway_ref, amount, status, paid_at", "gateway_ref UNIQUE = idempotency anchor"),
]
for i, (a, b, c) in enumerate(new_rows, start=1):
    r = t34.rows[i]
    r.cells[0].text, r.cells[1].text, r.cells[2].text = a, b, c
t34._tbl.remove(t34.rows[8]._tr)  # drop the old audit_log row

# ---- 7.3 UI notes ----
set_para(202, "Layout: a sidebar that changes with role (customer links; Kitchen for staff; Admin Dashboard, Menu & Stock, Time Slots, Users & Staff, All Orders and Reports for admins), menu cards with price, stock badges and instant search, a cart page with a slot dropdown showing seats left, an order page with timeline stepper, live badge, slot countdown and token box, kitchen tables with per-order countdowns and one-tap buttons, and a token handover box that upper-cases input by itself. JavaScript is vanilla: countdowns, order polling every 8 seconds, kitchen polling every 10, category helpers and form validation — no heavy framework, so even lab PCs stay fast.")

# ---- pseudocode touch-ups (structure kept for Ch.9) ----
sub_para(205, "qr=genQR(order.id)", "token=T…(order.id)")
sub_para(205, "QR exists IFF Paid; cash path does not exist.", "A token exists IFF Paid; a cash path does not exist.")
sub_para(205, "INSERT audit(Paid); COMMIT", "record Paid history; COMMIT")
sub_para(206, "create order(Placed, expires=now+15min)", "create order(PENDING_PAYMENT, expires=now+10min)")
sub_para(206, "mark Cancelled(expired), release hold, INSERT audit; COMMIT", "mark Failed, release the slot booking, record history; COMMIT")
sub_para(206, "NEVER generate QR before Paid; NEVER accept cash at any step.", "NEVER generate a token before Paid; NEVER accept cash at any step.")
sub_para(207, "schema isolates truth in eight tables;", "schema isolates truth in seven tables;")

# ---- Ch.8 intros ----
sub_para(210, "models, slot-capped checkout with 15-minute hold, idempotent online-only payment confirmation, QR gating, and KDS transitions.",
         "models, slot-capped checkout with 10-minute unpaid expiry, idempotent sandbox payment confirmation, token gating, and kitchen transitions.")
sub_para(211, "models.py (8 tables)", "models split by entity (7 tables)")
sub_para(211, "services.py (OrderService + PaymentService), pay_routes.py (create-order + webhook)",
         "services (order_service + payment_service), payments routes (checkout + webhook)")
sub_para(213, "Seven statuses, paise integers, unique transaction IDs, and hashed passwords. The CHECK that Paid implies a payment row is enforced in service logic plus a DB-level convention documented here; MySQL 8 CHECK syntax is included as a comment for Workbench forward-engineering.",
         "Eight statuses, DECIMAL rupees, unique gateway references, and scrypt-hashed passwords. The rule that Paid implies a payment row plus a token lives in the service layer (order_service + payment_service) — which is exactly what Chapter 9 tests.")

# ---- real model listing ----
set_para(214, "# app/models/user.py — passwords never touch the database in plain text\n"
"class User(UserMixin, db.Model):\n"
"    __tablename__ = \"users\"\n"
"    id = db.Column(db.Integer, primary_key=True)\n"
"    name = db.Column(db.String(80), nullable=False)\n"
"    email = db.Column(db.String(120), unique=True, nullable=False, index=True)\n"
"    password_hash = db.Column(db.String(255), nullable=False)\n"
"    role = db.Column(db.String(10), nullable=False, default=\"customer\")\n"
"    def set_password(self, raw):\n"
"        self.password_hash = generate_password_hash(raw)   # Werkzeug scrypt\n"
"    def check_password(self, raw):\n"
"        return check_password_hash(self.password_hash, raw)\n"
"# app/models/menu.py, category.py, slot.py (condensed)\n"
"class MenuItem(db.Model):  # menu_items: id, name, category, price NUMERIC(8,2), stock, is_available\n"
"class Category(db.Model):  # categories: id, name UNIQUE — addable without any item\n"
"class TimeSlot(db.Model):  # time_slots: id, start_time, capacity, booked — daily times, no dates")
set_para(215, "# app/models/order.py + payment.py (condensed)\n"
"class OrderStatus(enum.Enum):\n"
"    PENDING_PAYMENT, PAID, PREPARING, READY, COLLECTED, CANCELLED, FAILED, NO_SHOW\n"
"class Order(db.Model):  # orders: id, user_id FK, slot_id FK, total NUMERIC(10,2),\n"
"    # status ENUM(OrderStatus), token VARCHAR(10) UNIQUE — set ONLY after payment\n"
"class OrderItem(db.Model):  # order_items: id, order FK, item FK, quantity, unit_price\n"
"class Payment(db.Model):  # payments: id, order FK, gateway_ref, amount, status, paid_at")
set_para(216, "")
sub_para(217, "Listing 8.1 — Models: paise, enums, UNIQUE txn_id, QR-nullable-until-Paid.",
         "Listing 8.1 — Real models: scrypt passwords, eight statuses, DECIMAL totals, token-set-only-on-Paid.")

# ---- real service listing ----
set_para(220, "# app/services/order_service.py — place_order (condensed, comments are the test paths)\n"
"def place_order(user, cart, slot_id):\n"
"    if not cart: return fail(\"EMPTY_CART\")                                # P1\n"
"    slot = db.session.get(TimeSlot, slot_id)\n"
"    if slot is None: return fail(\"INVALID_SLOT\")                          # P2\n"
"    if slot.booked >= slot.capacity: return fail(\"SLOT_FULL\")             # P3\n"
"    if slot.start_time < utcnow() + 25 min: return fail(\"SLOT_TOO_SOON\")  # kitchen lead\n"
"    ...validate every line (available + stock)...                           # P4/P5\n"
"    order = create_order(...)  # PENDING_PAYMENT, slot booked\n"
"    payment = gateway.create_payment(order.id, total)  # sandbox adapter\n"
"    if not payment.ok: ... return fail(\"GATEWAY_ERROR\")                   # P6\n"
"    record Payment(...CREATED...); return ok(order_id, pay_url, total)")
set_para(221, "# app/services/payment_service.py — handle_event (condensed)\n"
"def handle_event(ref, success):\n"
"    payment = Payment.query.filter_by(gateway_ref=ref).first()\n"
"    if payment is None: return fail(\"UNKNOWN_PAYMENT\")\n"
"    if payment.status != CREATED: return ok(duplicate=True)  # webhook twice → ignore\n"
"    if success:\n"
"        payment.status = SUCCESS; order.status = PAID\n"
"        order.token = f\"T{order.id:04d}\"   # token exists IFF Paid\n"
"        stock decremented here — exactly once\n"
"    else: FAILED + slot booking released\n"
"    db.session.commit()\n"
"# Cancelling: PENDING_PAYMENT always free; PAID only within 5 minutes (then refund + stock back).")
sub_para(222, "Listing 8.2 — Hold + idempotent confirm. No cash path exists by construction.",
         "Listing 8.2 — Unpaid order + idempotent confirm. No cash path exists by construction.")
sub_para(219, "QR requires Paid", "a token requires Paid")

# ---- run instructions + takeaway ----
set_para(223, "Running it (lab or laptop): python -m venv venv; pip install -r requirements.txt; copy .env.example to .env; python seed.py (demo logins admin/staff/student @college.edu, password Password@123); python run.py → http://127.0.0.1:5000; pytest -q (19 tests, SQLite in-memory). Templates use Bootstrap 5 cards, a slot dropdown with seats-left badges, a Pay Online button, and a tracker with timeline plus countdown; the kitchen polls /staff/queue-data every 10 s. Hosting the same code: point DATABASE_URL at Supabase Postgres on Render with gunicorn run:app — tables and demo data build themselves on first boot. Screenshots of menu, cart, tracker and kitchen belong in Appendix D alongside the EER export.")
sub_para(224, "holds expire, payments verify, duplicates ignore, QR gates on Paid.",
         "unpaid orders lapse, payments verify, duplicates are ignored, tokens gate on Paid.")

d.save(P)
print("Batch C1 saved.")
